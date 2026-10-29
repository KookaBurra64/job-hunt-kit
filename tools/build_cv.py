import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# usage: python tools/build_cv.py OUT.docx [SOURCE.md]
# SOURCE.md defaults to my-cv.md. Line 1 is "# Your Name", line 2 is your contact line.
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if len(sys.argv) < 2:
    sys.exit('usage: python tools/build_cv.py output/My CV.docx [my-cv.md]')
src = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'my-cv.md')
md = open(src, encoding='utf-8').read().splitlines()
doc = Document()
# Tight margins keep a one-page CV on one page.
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.2); s.left_margin = s.right_margin = Cm(1.6)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(2); st.paragraph_format.space_before = Pt(0)

def add_runs(p, text):
    # **bold** segments
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)

after_name = False
for line in md:
    if not line.strip():
        continue
    contact = after_name and not line.startswith(('#', '- '))
    after_name = line.startswith('# ')
    if line.startswith('# '):
        p = doc.add_paragraph(); r = p.add_run(line[2:]); r.bold = True; r.font.size = Pt(20)
        p.paragraph_format.space_after = Pt(0)
    elif line.startswith('## '):
        p = doc.add_paragraph(); r = p.add_run(line[3:].upper()); r.bold = True; r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x33,0x33,0x33)
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        # rule
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
        b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:space'),'1'); b.set(qn('w:color'),'999999')
        pbdr.append(b); pPr.append(pbdr)
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); add_runs(p, line[2:])
        p.paragraph_format.space_after = Pt(1)
    else:
        p = doc.add_paragraph(); add_runs(p, line)
        if contact:
            p.runs[0].font.size = Pt(9.5); p.paragraph_format.space_after = Pt(6)

out = sys.argv[1]
os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
doc.save(out); print('saved', out, 'from', src)
# PDF next to the docx and a page count. Exit 1 if over the page limit (my-profile.toml, default 2).
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import myprofile
import render
sys.exit(render.report(out, max_pages=int(myprofile.load(required=False).get('cv', {}).get('max_pages', 2))))
