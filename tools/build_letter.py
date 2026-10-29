import os, sys
from docx import Document
from docx.shared import Pt, Cm

# usage: python tools/build_letter.py OUT.docx SOURCE.txt
# Name and contact line come from my-profile.toml, then one paragraph per blank-line block.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import myprofile
if len(sys.argv) < 3:
    sys.exit('usage: python tools/build_letter.py "output/Cover Letter - Company.docx" letters/company.txt')
me = myprofile.load().get('me', {})
out, src = sys.argv[1], sys.argv[2]
blocks = [b.strip() for b in open(src, encoding='utf-8').read().split('\n\n') if b.strip()]
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(8); st.paragraph_format.space_before = Pt(0)
p = doc.add_paragraph(); r = p.add_run(me.get('name', '')); r.bold = True; r.font.size = Pt(14)
p.paragraph_format.space_after = Pt(0)
doc.add_paragraph(' | '.join(v for v in (me.get(k, '') for k in ('location', 'phone', 'email', 'linkedin')) if v))
for b in blocks:
    if b.startswith('## '):          # selection criteria heading, for government-style letters
        h = doc.add_paragraph(); h.add_run(b[3:].strip()).bold = True
        h.paragraph_format.space_after = Pt(2); h.paragraph_format.keep_with_next = True
    else:
        doc.add_paragraph(b.replace('\n', ' '))
os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
doc.save(out); print('saved', out, 'from', src, f'({len(blocks)} paragraphs)')
# PDF next to the docx and a page count. Letters may run to two pages.
import render
sys.exit(render.report(out, max_pages=2))
